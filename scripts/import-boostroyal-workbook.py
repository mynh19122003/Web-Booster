"""Extract reference quotes without modifying the supplied workbook."""
import json
import sys
from pathlib import Path
import openpyxl

input_path = Path(sys.argv[1])
book = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
records = []
for row in book["Data_all"].iter_rows(min_row=5, values_only=True):
    if not isinstance(row[9], (int, float)):
        continue
    records.append({"game": row[1], "server": row[2], "from": row[3], "to": row[4], "gain": row[5], "currentPoints": row[6], "mode": row[7], "queue": row[8], "price": row[9], "originalPrice": row[10], "currency": row[12], "source": row[15]})
missing = [dict(zip(["server", "from", "to", "gain", "status"], row)) for row in book["Missing_LoL"].iter_rows(min_row=5, values_only=True) if row[0]]
output = Path(__file__).resolve().parents[1] / "data" / "pricing" / "boostroyal-workbook.json"
output.parent.mkdir(parents=True, exist_ok=True)
output.write_text(json.dumps({"sourceFile": input_path.name, "asOf": "2026-10-06", "quotes": records, "missing": missing}, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
print(f"Imported {len(records)} reference quotes and {len(missing)} missing combinations")
